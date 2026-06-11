from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "lab4_report.docx"
ASSETS = ROOT / "docs" / "assets"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )

    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")

    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    r_pr.append(color)

    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)

    new_run.append(r_pr)
    text_element = OxmlElement("w:t")
    text_element.text = text
    new_run.append(text_element)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def style_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for style_name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 16, 8),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ]:
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def add_title(doc):
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(3)
    run = title.add_run("Звіт до лабораторної роботи №4")
    run.bold = True
    run.font.name = "Calibri"
    run.font.size = Pt(22)
    run.font.color.rgb = RGBColor.from_string("0B2545")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(12)
    run = subtitle.add_run("Безперервна інтеграція та автоматизація розгортання (CI/CD)")
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor.from_string("555555")


def add_metadata(doc):
    table = doc.add_table(rows=4, cols=2)
    table.style = "Table Grid"
    table.autofit = False
    rows = [
        ("Проєкт", "UniDone"),
        ("Репозиторій", "https://github.com/wortpool/lab-4"),
        ("Production URL", "https://app-six-xi-33.vercel.app"),
        ("Workflow", "CI/CD Pipeline / build-and-test"),
    ]

    for row, (label, value) in zip(table.rows, rows):
        set_cell_width(row.cells[0], 2200)
        set_cell_width(row.cells[1], 6920)
        set_cell_shading(row.cells[0], "F2F4F7")
        row.cells[0].paragraphs[0].add_run(label).bold = True
        paragraph = row.cells[1].paragraphs[0]
        if value.startswith("https://"):
            add_hyperlink(paragraph, value, value)
        else:
            paragraph.add_run(value)

    doc.add_paragraph()


def add_labeled_paragraph(doc, label, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.add_run(label).bold = True
    paragraph.add_run(text)


def add_figure(doc, image_name, caption, width=6.25):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run().add_picture(str(ASSETS / image_name), width=Inches(width))

    caption_p = doc.add_paragraph()
    caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_p.paragraph_format.space_after = Pt(10)
    run = caption_p.add_run(caption)
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string("555555")


def build_docx():
    doc = Document()
    style_document(doc)
    add_title(doc)

    doc.add_paragraph(
        "Мета роботи: налаштувати pipeline as code для frontend MVP UniDone, "
        "автоматизувати перевірку якості коду, запуск unit-тестів, production build "
        "та розгортання застосунку у Vercel."
    )

    add_metadata(doc)

    doc.add_heading("1. GitHub Actions", level=1)
    doc.add_paragraph(
        "На вкладці Actions видно успішні запуски workflow. Останній запуск "
        "CI/CD Pipeline завершився зі статусом Success."
    )
    add_figure(
        doc,
        "actions-success.png",
        "Рисунок 1 - вкладка Actions у GitHub з успішною історією запусків.",
    )

    doc.add_heading("2. YAML workflow з поясненням", level=1)
    doc.add_paragraph(
        "Файл .github/workflows/main.yml описує CI/CD pipeline для проєкту UniDone."
    )
    add_labeled_paragraph(doc, "on.push.branches: ", "запускає pipeline при push у main та develop.")
    add_labeled_paragraph(doc, "pull_request: ", "запускає pipeline для перевірки змін перед merge.")
    add_labeled_paragraph(doc, "actions/checkout@v4: ", "завантажує код репозиторію на runner.")
    add_labeled_paragraph(doc, "actions/setup-node@v4: ", "встановлює Node.js 20 і вмикає npm cache.")
    add_labeled_paragraph(doc, "npm ci: ", "встановлює залежності з package-lock.json.")
    add_labeled_paragraph(doc, "npm run lint: ", "перевіряє якість коду через ESLint.")
    add_labeled_paragraph(doc, "npm run test:unit: ", "запускає unit-тести логіки задач.")
    add_labeled_paragraph(doc, "npm run build: ", "перевіряє production build Vite.")
    add_figure(
        doc,
        "workflow-yaml.png",
        "Рисунок 2 - YAML-файл GitHub Actions workflow.",
    )

    doc.add_heading("3. Репозиторій з бейджем у README.md", level=1)
    p = doc.add_paragraph("Репозиторій GitHub: ")
    add_hyperlink(p, "https://github.com/wortpool/lab-4", "https://github.com/wortpool/lab-4")
    doc.add_paragraph(
        "У верхній частині README.md додано status badge для workflow CI/CD Pipeline. "
        "Бейдж веде на сторінку GitHub Actions і показує актуальний стан перевірок."
    )

    doc.add_heading("4. Робочий веб-сайт UniDone", level=1)
    p = doc.add_paragraph("Production URL: ")
    add_hyperlink(p, "https://app-six-xi-33.vercel.app", "https://app-six-xi-33.vercel.app")
    doc.add_paragraph(
        "Сайт розгорнуто у Vercel. У production-версії застосунок відкривається без логіну "
        "і показує режим Mode: Production."
    )
    add_figure(
        doc,
        "production-site.png",
        "Рисунок 3 - production-версія UniDone у мережі Інтернет.",
    )

    doc.add_heading("Висновок", level=1)
    doc.add_paragraph(
        "Для проєкту UniDone налаштовано CI/CD workflow, автоматизовано перевірку якості, "
        "unit-тести та production build. Репозиторій підключено до Vercel, а готовий сайт "
        "розміщено у мережі Інтернет."
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)


if __name__ == "__main__":
    build_docx()
    print(OUT)
